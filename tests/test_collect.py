import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from collect import (
    classify_job,
    extract,
    extract_attachment_text,
    iso_date,
    is_recruitment_notice,
    merge_job,
    normalize_title,
    parse_detail_html,
    run_collection,
)

SOURCE = {
    "name": "测试医院",
    "type": "医院",
    "website": "https://jobs.example.org/",
    "url": "https://jobs.example.org/list",
    "domain": "jobs.example.org",
    "method": "测试 HTML",
    "detail_delay": 0,
    "list_delay": 0,
}
STAMP = "2026-10-08T09:25:00+08:00"


class FakeResponse:
    def __init__(self, text="", status=200, content=b"", headers=None):
        self.text = text
        self.status_code = status
        self.encoding = "utf-8"
        self.apparent_encoding = "utf-8"
        self.content = content
        self.headers = headers or {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code} response")

    def iter_content(self, _size):
        yield self.content


class FakeSession:
    def __init__(self, routes):
        self.routes = routes

    def get(self, url, **_kwargs):
        result = self.routes[url]
        if isinstance(result, Exception):
            raise result
        return result


class CollectTests(unittest.TestCase):
    def test_date_parsing_and_invalid_date(self):
        self.assertEqual(iso_date("发布日期：2026年09月28日"), "2026-09-28")
        self.assertEqual(iso_date("2026/9/8"), "2026-09-08")
        self.assertIsNone(iso_date("2026-02-31"))
        self.assertIsNone(iso_date("未核实"))

    def test_title_date_is_removed_and_listing_date_is_kept(self):
        html = """<ul><li><span>2026-09-28</span>
        <a href="/a">2026-09-28 山东医院博士人才招聘公告</a></li></ul>"""
        item = extract(html, SOURCE, STAMP)[0]
        self.assertEqual(item["title"], "山东医院博士人才招聘公告")
        self.assertEqual(item["published"], "2026-09-28")
        self.assertEqual(normalize_title("【2026年9月28日】 招聘公告"), "招聘公告")

    def test_extract_deduplicates_and_filters_non_recruitment(self):
        html = """<div>
          <a href="/a">医院博士人才招聘公告</a>
          <a href="/a#top">医院博士人才招聘公告</a>
          <a href="/b">公开招聘拟聘人员公示</a>
          <a href="/c">住院医师规范化培训招收简章</a>
          <a href="/d">博士研究生招生简章</a>
          <a href="https://bad.example/x">医院招聘公告信息</a>
        </div>"""
        items = extract(html, SOURCE, STAMP)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["url"], "https://jobs.example.org/a")
        self.assertFalse(is_recruitment_notice("公开招聘考试安排"))

    def test_detail_dates_fields_and_attachment_links(self):
        html = """<html><head><meta name="publishdate" content="2026-09-20"></head><body>
        <p>报名时间：2026年9月21日至2026年10月8日。</p>
        <p>学历要求：博士研究生。学位要求：博士学位。</p>
        <p>招聘人数：3人。工作地点：山东省济南市历下区。</p>
        <p>基础医学、药理学、分子生物学方向优先。</p>
        <a href="/files/positions.xlsx">岗位表</a>
        <a href="https://other.example/file.pdf">站外文件</a>
        </body></html>"""
        detail = parse_detail_html(html, "https://jobs.example.org/a", SOURCE)
        self.assertEqual(detail["published"], "2026-09-20")
        self.assertEqual(detail["application_start"], "2026-09-21")
        self.assertEqual(detail["deadline"], "2026-10-08")
        self.assertEqual(detail["recruitment_count"], "3人")
        self.assertIn("基础医学", detail["specialties"])
        self.assertEqual(len(detail["attachments"]), 1)
        self.assertIn("deadline", detail["evidence"])

    def test_phd_specialty_scoring_is_explainable(self):
        job = {"title": "医院2026年高层次人才招聘公告"}
        result = classify_job(job, "招聘基础医学博士，从事专职科研工作")
        self.assertEqual(result["category"], "博士岗位")
        self.assertGreaterEqual(result["score"], 6)
        self.assertIn("基础医学", result["specialties"])
        self.assertTrue(result["reasons"])

    def test_clinical_restriction_requires_manual_review(self):
        job = {"title": "肾内科工作人员招聘公告"}
        result = classify_job(job, "专业为肾脏病学，须取得执业医师证和住培合格证")
        self.assertEqual(result["category"], "需要人工核实")
        self.assertIn("要求临床执业资格", result["restrictions"])
        self.assertIn("要求住院医师规范化培训", result["restrictions"])

    def test_docx_and_xlsx_attachments_are_parsed(self):
        from docx import Document
        from openpyxl import Workbook

        doc = Document()
        doc.add_paragraph("博士岗位 基础医学")
        doc_buffer = io.BytesIO()
        doc.save(doc_buffer)
        text, error = extract_attachment_text(doc_buffer.getvalue(), "岗位说明.docx")
        self.assertIsNone(error)
        self.assertIn("基础医学", text)

        workbook = Workbook()
        workbook.active.append(["岗位", "专业"])
        workbook.active.append(["科研岗", "药理学"])
        xlsx_buffer = io.BytesIO()
        workbook.save(xlsx_buffer)
        text, error = extract_attachment_text(xlsx_buffer.getvalue(), "岗位表.xlsx")
        self.assertIsNone(error)
        self.assertIn("药理学", text)

    def test_merge_tracks_revision_and_preserves_verified_deadline(self):
        current = {
            "id": "a", "title": "博士招聘公告（修订）", "published": None,
            "deadline": None, "education": "未核实", "degree": "未核实",
            "specialties": [], "recruitment_count": "未核实", "location": "未核实",
            "attachments": [], "phd_relevance": {}, "restrictions": [],
        }
        previous = {
            "id": "a", "title": "博士招聘公告", "published": "2026-09-01",
            "published_evidence": "列表项", "deadline": "2026-10-01", "education": "博士研究生",
            "degree": "博士", "specialties": [], "recruitment_count": "未核实",
            "location": "未核实", "attachments": [], "phd_relevance": {}, "restrictions": [],
            "first_seen": "2026-09-01T09:00:00+08:00", "revisions": [],
        }
        merged = merge_job(current, previous, STAMP)
        self.assertEqual(merged["deadline"], "2026-10-01")
        self.assertEqual(merged["published"], "2026-09-01")
        self.assertEqual(len(merged["revisions"]), 1)

    @patch("collect.time.sleep", return_value=None)
    def test_source_failure_retains_history_and_unknown_count(self, _sleep):
        previous = {
            "sources": [{"name": "测试医院", "last_success_at": "2026-10-07T09:25:00+08:00"}],
            "jobs": [{"id": "old", "source": "测试医院", "title": "历史博士招聘公告",
                      "url": "https://jobs.example.org/old", "first_seen": "2026-10-01T09:00:00+08:00"}],
        }
        session = FakeSession({SOURCE["url"]: requests.ConnectionError("proxy denied")})
        document, exit_code = run_collection([SOURCE], previous, session, STAMP)
        self.assertEqual(exit_code, 1)
        self.assertEqual(len(document["jobs"]), 1)
        self.assertEqual(document["sources"][0]["status"], "error")
        self.assertIsNone(document["sources"][0]["found"])
        self.assertEqual(document["sources"][0]["last_success_at"], "2026-10-07T09:25:00+08:00")

    @patch("collect.time.sleep", return_value=None)
    def test_json_schema_and_successful_collection(self, _sleep):
        listing = '<li>2026-10-07 <a href="/a">医院博士科研岗位招聘公告</a></li>'
        detail = "<p>报名截止时间：2026年10月20日。基础医学博士。</p>"
        session = FakeSession({
            SOURCE["url"]: FakeResponse(listing),
            "https://jobs.example.org/a": FakeResponse(detail),
        })
        document, exit_code = run_collection([SOURCE], {}, session, STAMP)
        self.assertEqual(exit_code, 0)
        self.assertEqual(document["schema_version"], 2)
        self.assertEqual(document["collection_status"], "ok")
        self.assertEqual(document["statistics"]["jobs_total"], 1)
        self.assertEqual(document["jobs"][0]["phd_relevance"]["category"], "博士岗位")
        json.dumps(document, ensure_ascii=False)


if __name__ == "__main__":
    unittest.main()
