#!/usr/bin/env python3
"""Collect and classify official recruitment notices without bypassing site controls."""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import sys
import time
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urldefrag, urljoin, urlparse, urlunparse

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

ROOT = Path(__file__).resolve().parents[1]
SOURCES = json.loads((ROOT / "scripts/sources.json").read_text(encoding="utf-8"))
OUTPUT = ROOT / "site/data/jobs.json"
LOCAL = timezone(timedelta(hours=8))
HEADERS = {
    "User-Agent": "JinanJobRadar/1.0 (+https://github.com/dujunyao044818-web/jinan-job-radar; public recruitment monitor)",
    "Accept-Language": "zh-CN,zh;q=0.9",
}
MAX_ATTACHMENT_BYTES = 12 * 1024 * 1024

DATE = re.compile(r"(?<!\d)(20\d{2})\s*[年./-]\s*(0?[1-9]|1[0-2])\s*[月./-]\s*(3[01]|[12]\d|0?[1-9])\s*日?")
DATE_PREFIX = re.compile(
    r"^\s*(?:\[|【|（|\()?\s*20\d{2}\s*[年./-]\s*\d{1,2}\s*[月./-]\s*\d{1,2}\s*日?\s*(?:\]|】|）|\))?\s*[-—·|:]?\s*"
)
POSITIVE = re.compile(r"招聘|招募|招收|引进|聘用|诚聘|博士后")
NEGATIVE = re.compile(
    r"拟聘|公示|名单|面试|笔试|成绩|资格复审|资格审查|体检|录取|考察|考核范围|"
    r"考试安排|核减|取消岗位|补录|递补|录用结果|招聘结果|进修|培训班|住院医师规范化培训|规培|"
    r"研究生招生|博士研究生招生|招生简章|夏令营|宣讲会"
)
SPECIALTY_PATTERNS = {
    "基础医学": re.compile(r"基础医学|病理生理|人体解剖|组织胚胎"),
    "药理学": re.compile(r"药理学|药物化学|药剂学|药学"),
    "肾脏病学": re.compile(r"肾脏病|肾病|肾内"),
    "临床检验诊断学": re.compile(r"临床检验诊断|检验诊断"),
    "医学检验": re.compile(r"医学检验|检验技术|检验科"),
    "生物医学": re.compile(r"生物医学|生物工程"),
    "分子生物学": re.compile(r"分子生物|细胞生物|生物化学"),
}
RESTRICTION_PATTERNS = {
    "要求医师资格证": re.compile(r"医师资格证|医师资格考试|执业医师资格"),
    "要求住院医师规范化培训": re.compile(r"住院医师规范化培训|住培合格|规培合格"),
    "要求临床执业资格": re.compile(r"临床执业资格|执业医师证|医师执业证"),
    "要求特定职称": re.compile(r"主任医师|副主任医师|高级职称|中级职称"),
}
DEADLINE_PATTERNS = [
    re.compile(r"(?:报名|申请|申报)(?:时间|期限|截止(?:时间|日期)?)?[^。；\n]{0,100}?(?:至|截止(?:到|至)?|结束于)\s*(20\d{2}[年./-]\d{1,2}[月./-]\d{1,2}日?)"),
    re.compile(r"(?:报名|申请|申报)截止(?:时间|日期)?\s*[：:]?\s*(20\d{2}[年./-]\d{1,2}[月./-]\d{1,2}日?)"),
]
START_PATTERNS = [
    re.compile(r"(?:报名|申请|申报)(?:开始)?时间\s*[：:]?\s*(?:自)?\s*(20\d{2}[年./-]\d{1,2}[月./-]\d{1,2}日?)"),
    re.compile(r"(?:报名|申请|申报)\s*自\s*(20\d{2}[年./-]\d{1,2}[月./-]\d{1,2}日?)\s*(?:起|开始)"),
]


def now_iso() -> str:
    return datetime.now(LOCAL).isoformat(timespec="seconds")


def iso_date(text: str | None) -> str | None:
    match = DATE.search(text or "")
    if not match:
        return None
    try:
        return datetime(*(int(part) for part in match.groups())).date().isoformat()
    except ValueError:
        return None


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_title(value: str | None) -> str:
    title = clean_text(value).strip("·-—>| ")
    previous = None
    while title and previous != title:
        previous = title
        title = DATE_PREFIX.sub("", title).strip("·-—>| ")
    return title


def is_recruitment_notice(title: str) -> bool:
    return 8 <= len(title) <= 180 and bool(POSITIVE.search(title)) and not bool(NEGATIVE.search(title))


def canonical_url(url: str) -> str:
    url = urldefrag(url)[0]
    parsed = urlparse(url)
    query = [(k, v) for k, v in parse_qsl(parsed.query, keep_blank_values=True) if not k.lower().startswith(("utm_", "spm"))]
    path = re.sub(r"/{2,}", "/", parsed.path or "/")
    return urlunparse((parsed.scheme.lower(), parsed.netloc.lower(), path, "", urlencode(query), ""))


def valid_link(url: str, source: dict[str, Any]) -> bool:
    parsed = urlparse(url)
    allowed = set(source.get("allowed_domains") or [source["domain"]])
    blocked = re.search(r"\.(?:jpg|jpeg|png|gif|css|js|pdf|docx?|xlsx?|zip|rar)(?:$|\?)", parsed.path, re.I)
    return parsed.scheme in {"http", "https"} and parsed.hostname in allowed and not blocked


def listing_context(anchor: Any) -> str:
    for parent in anchor.parents:
        if getattr(parent, "name", None) in {"li", "tr", "article", "section", "dd", "dt"}:
            return clean_text(parent.get_text(" ", strip=True))[:500]
    return clean_text(anchor.parent.get_text(" ", strip=True) if anchor.parent else anchor.get_text(" ", strip=True))[:300]


def extract(html: str, source: dict[str, Any], collected_at: str | None = None) -> list[dict[str, Any]]:
    """Extract candidate notices from one official listing page."""
    soup = BeautifulSoup(html, "html.parser")
    seen: dict[str, dict[str, Any]] = {}
    timestamp = collected_at or now_iso()
    for anchor in soup.select(source.get("link_selector", "a[href]")):
        raw_title = anchor.get("title") or anchor.get_text(" ", strip=True)
        title = normalize_title(raw_title)
        if not is_recruitment_notice(title):
            continue
        url = canonical_url(urljoin(source["url"], anchor.get("href", "")))
        if not valid_link(url, source):
            continue
        context = listing_context(anchor)
        published = iso_date(context.replace(clean_text(raw_title), " ", 1)) or iso_date(raw_title)
        key = hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]
        seen[key] = {
            "id": key, "source": source["name"], "source_aliases": source.get("aliases", []),
            "type": source["type"], "title": title, "url": url, "published": published,
            "published_evidence": f"列表项：{context[:180]}" if published else None,
            "application_start": None, "deadline": None, "deadline_status": "未核实",
            "deadline_reliable": False,
            "education": "未核实", "degree": "未核实", "specialties": [],
            "recruitment_count": "未核实", "location": "未核实", "attachments": [],
            "evidence": {}, "collected_at": timestamp,
        }
    return list(seen.values())


def build_session() -> requests.Session:
    retry = Retry(
        total=2, connect=2, read=2, backoff_factor=0.8,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET", "HEAD"}), respect_retry_after_header=True,
    )
    session = requests.Session()
    session.headers.update(HEADERS)
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.mount("http://", HTTPAdapter(max_retries=retry))
    ca_bundle = os.environ.get("REQUESTS_CA_BUNDLE") or os.environ.get("SSL_CERT_FILE")
    if ca_bundle and Path(ca_bundle).is_file():
        session.verify = ca_bundle
    return session


def _match_first(patterns: list[re.Pattern[str]], text: str) -> tuple[str | None, str | None]:
    for pattern in patterns:
        match = pattern.search(text)
        if match:
            value = iso_date(match.group(1))
            if value:
                return value, clean_text(match.group(0))[:180]
    return None, None


def extract_attachment_text(content: bytes, filename: str) -> tuple[str | None, str | None]:
    """Parse PDF, DOCX and XLSX; retain links to unsupported or damaged attachments."""
    suffix = Path(urlparse(filename).path).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages), None
        if suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            return "\n".join(parts), None
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(sheet.title)
                rows.extend(" | ".join(clean_text(str(cell)) for cell in row if cell is not None) for row in sheet.iter_rows(values_only=True))
            return "\n".join(rows), None
        if suffix in {".doc", ".xls"}:
            return None, "旧版二进制 Office 格式暂不自动解析，请打开官方附件核实"
        return None, "不支持的附件格式，请打开官方附件核实"
    except Exception as exc:
        return None, f"附件解析失败：{type(exc).__name__}"


def discover_attachments(soup: BeautifulSoup, base_url: str, source: dict[str, Any]) -> list[dict[str, Any]]:
    attachments = []
    allowed = set(source.get("allowed_domains") or [source["domain"]])
    for anchor in soup.select("a[href]"):
        url = canonical_url(urljoin(base_url, anchor.get("href", "")))
        parsed = urlparse(url)
        suffix = Path(parsed.path).suffix.lower()
        if parsed.hostname not in allowed or suffix not in {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".zip", ".rar"}:
            continue
        attachments.append({
            "name": clean_text(anchor.get_text(" ", strip=True)) or Path(parsed.path).name,
            "url": url, "format": suffix.lstrip("."), "parsed": False, "parse_error": None,
        })
    return list({item["url"]: item for item in attachments}.values())


def parse_detail_html(html: str, url: str, source: dict[str, Any]) -> dict[str, Any]:
    soup = BeautifulSoup(html, "html.parser")
    attachments = discover_attachments(soup, url, source)
    published = published_evidence = None
    for selector in ("meta[name='PubDate']", "meta[name='publishdate']", "meta[property='article:published_time']"):
        tag = soup.select_one(selector)
        if tag and iso_date(tag.get("content")):
            published = iso_date(tag.get("content"))
            published_evidence = f"页面元数据：{tag.get('content')}"
            break
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    text = clean_text(soup.get_text(" ", strip=True))
    deadline, deadline_evidence = _match_first(DEADLINE_PATTERNS, text)
    start, start_evidence = _match_first(START_PATTERNS, text)
    specialties = [name for name, pattern in SPECIALTY_PATTERNS.items() if pattern.search(text)]
    education_match = re.search(r"(?:学历|学历要求)[：:]?\s*([^，。；]{2,30})", text)
    degree_match = re.search(r"(?:学位|学位要求)[：:]?\s*([^，。；]{2,30})", text)
    count_match = re.search(r"(?:招聘|需求|计划)(?:人数|数量)[：:]?\s*(\d+\s*人?)", text)
    location_match = re.search(r"(?:工作地点|工作地址|工作所在地)[：:]?\s*([^，。；]{2,50})", text)
    evidence = {
        "published": published_evidence, "application_start": start_evidence, "deadline": deadline_evidence,
        "education": clean_text(education_match.group(0))[:180] if education_match else None,
        "degree": clean_text(degree_match.group(0))[:180] if degree_match else None,
        "recruitment_count": clean_text(count_match.group(0))[:180] if count_match else None,
        "location": clean_text(location_match.group(0))[:180] if location_match else None,
    }
    return {
        "text": text, "published": published, "published_evidence": published_evidence,
        "application_start": start, "deadline": deadline,
        "deadline_status": "从官方公告正文提取，需核对原文" if deadline else "未核实",
        "deadline_reliable": bool(deadline),
        "specialties": specialties,
        "education": clean_text(education_match.group(1)) if education_match else "未核实",
        "degree": clean_text(degree_match.group(1)) if degree_match else "未核实",
        "recruitment_count": clean_text(count_match.group(1)) if count_match else "未核实",
        "location": clean_text(location_match.group(1)) if location_match else "未核实",
        "attachments": attachments, "evidence": {key: value for key, value in evidence.items() if value},
    }


def classify_job(job: dict[str, Any], detail_text: str = "", attachment_text: str = "") -> dict[str, Any]:
    combined = clean_text(" ".join((job.get("title", ""), detail_text, attachment_text)))
    score, reasons = 0, []
    if re.search(r"博士后", combined):
        score += 7
        reasons.append("公告明确提及博士后")
    elif re.search(r"博士|博士学位|高层次人才", combined):
        score += 6
        reasons.append("公告明确提及博士或高层次人才")
    elif re.search(r"专职科研|科研岗|教学科研", combined):
        score += 3
        reasons.append("公告提及科研或教学科研岗位")
    elif re.search(r"优秀人才|青年人才|领军人才", combined):
        score += 2
        reasons.append("公告提及人才岗位，但学历要求仍需核实")
    specialties = sorted({name for name, pattern in SPECIALTY_PATTERNS.items() if pattern.search(combined)})
    if specialties:
        score += 3
        reasons.append("匹配优先专业：" + "、".join(specialties))
    restrictions = [label for label, pattern in RESTRICTION_PATTERNS.items() if pattern.search(combined)]
    if restrictions:
        reasons.append("存在资格限制，需人工核实")
    if NEGATIVE.search(job.get("title", "")):
        category, score, reasons = "非目标岗位", 0, ["标题属于结果、公示、考试或招生信息"]
    elif score >= 6:
        category = "博士岗位"
    elif specialties and score >= 3:
        category = "可能符合专业"
    elif score > 0:
        category = "需要人工核实"
    else:
        category = "非目标岗位"
    if restrictions and specialties and not re.search(r"博士|科研岗|专职科研|教学科研", combined):
        category = "需要人工核实"
        reasons.append("专业匹配不能替代临床执业条件")
    return {
        "category": category, "score": score, "reasons": reasons or ["未识别到博士或优先专业依据"],
        "specialties": specialties, "restrictions": restrictions, "is_postdoc": bool(re.search(r"博士后", combined)),
    }


def fetch_attachment(session: requests.Session, attachment: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    try:
        response = session.get(attachment["url"], timeout=(8, 25), stream=True)
        response.raise_for_status()
        if int(response.headers.get("content-length", 0) or 0) > MAX_ATTACHMENT_BYTES:
            attachment["parse_error"] = "附件超过 12MB，保留链接供人工核实"
            return "", attachment
        chunks, size = [], 0
        for chunk in response.iter_content(64 * 1024):
            size += len(chunk)
            if size > MAX_ATTACHMENT_BYTES:
                attachment["parse_error"] = "附件超过 12MB，保留链接供人工核实"
                return "", attachment
            chunks.append(chunk)
        text, error = extract_attachment_text(b"".join(chunks), attachment["url"])
        attachment["parsed"], attachment["parse_error"] = text is not None, error
        return clean_text(text), attachment
    except requests.RequestException as exc:
        attachment["parse_error"] = f"附件下载失败：{type(exc).__name__}"
        return "", attachment


def enrich_job(job: dict[str, Any], source: dict[str, Any], session: requests.Session) -> dict[str, Any]:
    response = session.get(job["url"], timeout=(8, 25))
    response.raise_for_status()
    if response.encoding is None or response.encoding.lower() == "iso-8859-1":
        response.encoding = response.apparent_encoding or "utf-8"
    detail = parse_detail_html(response.text, job["url"], source)
    job["published"] = job.get("published") or detail["published"]
    job["published_evidence"] = job.get("published_evidence") or detail["published_evidence"]
    for field in ("application_start", "deadline", "deadline_status", "deadline_reliable", "education", "degree", "specialties", "recruitment_count", "location", "attachments", "evidence"):
        job[field] = detail[field]
    attachment_texts = []
    for attachment in job["attachments"][:8]:
        text, _ = fetch_attachment(session, attachment)
        if text:
            attachment_texts.append(text[:150_000])
        time.sleep(0.15)
    relevance = classify_job(job, detail["text"][:200_000], " ".join(attachment_texts))
    job["phd_relevance"], job["phd"] = relevance, relevance["category"] == "博士岗位"
    job["specialties"] = sorted(set(job["specialties"]) | set(relevance["specialties"]))
    job["restrictions"] = relevance["restrictions"]
    return job


def record_fingerprint(job: dict[str, Any]) -> str:
    keys = ("title", "published", "application_start", "deadline", "deadline_reliable", "education", "degree", "specialties", "recruitment_count", "location", "attachments", "phd_relevance", "restrictions")
    tracked = {key: job.get(key) for key in keys}
    return hashlib.sha256(json.dumps(tracked, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:20]


def merge_job(current: dict[str, Any], previous: dict[str, Any] | None, timestamp: str) -> dict[str, Any]:
    if not previous:
        current.update(first_seen=timestamp, last_seen=timestamp, is_newly_discovered=True, revisions=[])
        current["fingerprint"] = record_fingerprint(current)
        return current
    current.update(first_seen=previous.get("first_seen", timestamp), last_seen=timestamp, is_newly_discovered=False)
    current["published"] = current.get("published") or previous.get("published")
    current["published_evidence"] = current.get("published_evidence") or previous.get("published_evidence")
    for field in ("application_start", "deadline", "education", "degree", "recruitment_count", "location"):
        if current.get(field) in (None, "未核实", "") and previous.get(field) not in (None, "未核实", ""):
            current[field] = previous[field]
    if current.get("deadline") and current.get("deadline") == previous.get("deadline"):
        current["deadline_reliable"] = current.get("deadline_reliable") or previous.get("deadline_reliable", False)
    revisions = list(previous.get("revisions", []))[-9:]
    new_fingerprint = record_fingerprint(current)
    old_fingerprint = previous.get("fingerprint") or record_fingerprint(previous)
    if new_fingerprint != old_fingerprint:
        revisions.append({"observed_at": timestamp, "previous_fingerprint": old_fingerprint, "summary": "公告结构化字段发生变化，请核对官方原文"})
    current["revisions"], current["fingerprint"] = revisions, new_fingerprint
    return current


def collect_source(source: dict[str, Any], session: requests.Session, timestamp: str) -> list[dict[str, Any]]:
    response = session.get(source["url"], timeout=(8, 25))
    response.raise_for_status()
    if response.encoding is None or response.encoding.lower() == "iso-8859-1":
        response.encoding = response.apparent_encoding or "utf-8"
    jobs = extract(response.text, source, timestamp)
    if not jobs:
        raise ValueError("招聘栏目可访问，但未识别到目标招聘公告；可能是页面结构变化或当前无可验证条目")
    detail_limit = int(source.get("detail_limit", 12))
    for job in jobs[:detail_limit]:
        try:
            enrich_job(job, source, session)
        except (requests.RequestException, ValueError) as exc:
            relevance = classify_job(job)
            job.update(phd_relevance=relevance, phd=relevance["category"] == "博士岗位", restrictions=relevance["restrictions"], detail_error=f"详情解析失败：{type(exc).__name__}")
        time.sleep(float(source.get("detail_delay", 0.35)))
    for job in jobs[detail_limit:]:
        relevance = classify_job(job)
        job.update(phd_relevance=relevance, phd=relevance["category"] == "博士岗位", restrictions=relevance["restrictions"], detail_error="本轮详情解析预算已用完，需人工核实")
    return jobs


def load_previous(path: Path = OUTPUT) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def migrate_previous_jobs(previous: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Normalize legacy IDs and flag filtered records while retaining their history."""
    records: dict[str, dict[str, Any]] = {}
    for original in previous.get("jobs", []):
        if not original.get("url"):
            continue
        job = deepcopy(original)
        job["url"] = canonical_url(job["url"])
        job["id"] = hashlib.sha256(job["url"].encode("utf-8")).hexdigest()[:20]
        job["title"] = normalize_title(job.get("title"))
        relevance = job.get("phd_relevance") or classify_job(job)
        job["phd_relevance"] = relevance
        job["phd"] = relevance["category"] == "博士岗位"
        job.setdefault("specialties", relevance["specialties"])
        job.setdefault("restrictions", relevance["restrictions"])
        job["excluded"] = not is_recruitment_notice(job["title"])
        job["excluded_reason"] = "非招聘公告、结果公示、考试安排、规培或招生信息" if job["excluded"] else None
        existing = records.get(job["id"])
        if not existing or len(job) > len(existing):
            records[job["id"]] = job
    return records


def run_collection(sources=None, previous=None, session=None, timestamp=None) -> tuple[dict[str, Any], int]:
    sources = sources or [source for source in SOURCES if source.get("enabled", True)]
    previous, session, timestamp = previous or {}, session or build_session(), timestamp or now_iso()
    records = migrate_previous_jobs(previous)
    prior_statuses = {status.get("name"): status for status in previous.get("sources", [])}
    statuses, successful = [], 0
    for source in sources:
        prior = prior_statuses.get(source["name"], {})
        try:
            fresh = collect_source(source, session, timestamp)
            new_count = 0
            for job in fresh:
                before = records.get(job["id"])
                job["excluded"] = False
                job["excluded_reason"] = None
                records[job["id"]] = merge_job(job, before, timestamp)
                new_count += before is None
            successful += 1
            statuses.append({
                "name": source["name"], "type": source["type"], "website": source.get("website"),
                "url": source["url"], "method": source.get("method", "HTML 列表与详情解析"),
                "status": "ok", "found": len(fresh), "new": new_count,
                "last_attempt_at": timestamp, "last_success_at": timestamp, "error": "",
            })
        except (requests.RequestException, ValueError) as exc:
            statuses.append({
                "name": source["name"], "type": source["type"], "website": source.get("website"),
                "url": source["url"], "method": source.get("method", "HTML 列表与详情解析"),
                "status": "error", "found": None, "new": 0, "last_attempt_at": timestamp,
                "last_success_at": prior.get("last_success_at"), "error": clean_text(str(exc))[:240],
            })
        time.sleep(float(source.get("list_delay", 0.8)))
    jobs = sorted(records.values(), key=lambda item: (item.get("published") or "", item.get("first_seen") or ""), reverse=True)[:3000]
    document = {
        "schema_version": 2, "updated_at": timestamp,
        "collection_status": "ok" if successful == len(sources) else "partial" if successful else "failed",
        "statistics": {
            "sources_total": len(sources), "sources_successful": successful, "jobs_total": len(jobs),
            "phd_jobs": sum(job.get("phd_relevance", {}).get("category") == "博士岗位" or job.get("phd") is True for job in jobs),
        },
        "sources": statuses, "jobs": jobs,
        "disclaimer": "公告索引仅供参考；是否招聘、报名期限、学历专业和执业资格要求以官方原文与附件为准。未识别字段均为未核实。",
    }
    return document, 0 if successful else 1


def main() -> int:
    document, exit_code = run_collection(previous=load_previous())
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    stats = document["statistics"]
    print(f"Sources OK: {stats['sources_successful']} / {stats['sources_total']} | indexed: {stats['jobs_total']}")
    for status in document["sources"]:
        count = "未知" if status["found"] is None else status["found"]
        print(status["name"], status["status"], count, status["error"])
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
